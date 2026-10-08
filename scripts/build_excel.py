"""Build excel/HR_Attrition_Analysis.xlsx (formula-driven) from the clean CSV.

Usage (project root):  python scripts/build_excel.py
All analysis cells are live formulas (COUNTIFS / AVERAGE) over the Data sheet.
"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

ROOT = Path(__file__).resolve().parents[1]
df = pd.read_csv(ROOT / "data" / "hr_attrition_clean.csv")
n = len(df)
last = n + 1

F = "Arial"
H = Font(name=F, bold=True, color="FFFFFF", size=10)
HF = PatternFill("solid", fgColor="2B6CB0")
B = Font(name=F, size=10)
BB = Font(name=F, size=10, bold=True)
T = Font(name=F, size=14, bold=True)
NOTE = Font(name=F, size=9, italic=True, color="555555")
thin = Side(style="thin", color="D0D7DE")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()

# ---------------- Data ----------------
ws = wb.active
ws.title = "Data"
cols = list(df.columns)
ws.append(cols + ["age_band", "tenure_band"])
for row in df.itertuples(index=False):
    ws.append(list(row))
col = {c: L(i + 1) for i, c in enumerate(cols)}
ab, tb = L(len(cols) + 1), L(len(cols) + 2)
for r in range(2, last + 1):
    a, y = f"{col['age']}{r}", f"{col['years_at_company']}{r}"
    ws[f"{ab}{r}"] = f'=IF({a}<25,"Under 25",IF({a}<35,"25-34",IF({a}<45,"35-44",IF({a}<55,"45-54","55+"))))'
    ws[f"{tb}{r}"] = f'=IF({y}<=1,"0-1 yr",IF({y}<=3,"2-3 yrs",IF({y}<=5,"4-5 yrs",IF({y}<=10,"6-10 yrs","10+ yrs"))))'
for c in ws[1]:
    c.font, c.fill = H, HF
    c.alignment = Alignment(horizontal="center", wrap_text=True)
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.font = B
ws.freeze_panes = "A2"
for i in range(1, len(cols) + 3):
    ws.column_dimensions[L(i)].width = 17
ws.auto_filter.ref = f"A1:{tb}{last}"

def rng(name):                        # absolute range on Data sheet
    c = ab if name == "age_band" else tb if name == "tenure_band" else col[name]
    return f"Data!${c}$2:${c}${last}"

ATTR = rng("attrition")

# ---------------- KPIs ----------------
k = wb.create_sheet("KPIs")
k["A1"], k["A1"].font = "HR Attrition: headline KPIs", T
for i, h in enumerate(["Metric", "Value"], 1):
    c = k.cell(3, i, h); c.font, c.fill, c.border = H, HF, BOX
rows = [
    ("Employees", f"=COUNTA({rng('employee_id')})", "#,##0"),
    ("Attrition count (left)", f'=COUNTIF({ATTR},"Yes")', "#,##0"),
    ("Attrition rate", "=B5/B4", "0.00%"),
    ("Active employees", f'=COUNTIF({ATTR},"No")', "#,##0"),
    ("Average age", f"=AVERAGE({rng('age')})", "0.00"),
    ("Average monthly income (left)", f'=AVERAGEIF({ATTR},"Yes",{rng("monthly_income")})', "#,##0"),
    ("Average monthly income (stayed)", f'=AVERAGEIF({ATTR},"No",{rng("monthly_income")})', "#,##0"),
]
for i, (m, f, nf) in enumerate(rows, 4):
    a, b = k.cell(i, 1, m), k.cell(i, 2, f)
    a.font = BB; b.font = B; b.number_format = nf
    a.border = b.border = BOX
k["A12"] = "Source: data/hr_attrition_clean.csv (IBM HR Analytics sample data, fictional employees). All values are formulas over the Data sheet."
k["A12"].font = NOTE
k.column_dimensions["A"].width = 36; k.column_dimensions["B"].width = 16

# ---------------- Analysis ----------------
a = wb.create_sheet("Analysis")
a["A1"], a["A1"].font = "Attrition by segment (all formulas)", T
a["A2"] = "Age and tenure bands are formulas in Data!AG:AH. Cutoffs: age <25 / <35 / <45 / <55; tenure <=1 / <=3 / <=5 / <=10 years."
a["A2"].font = NOTE
a.column_dimensions["A"].width = 26
for c in "BCD":
    a.column_dimensions[c].width = 13

row_ptr = 4
blocks = {}

def block(title, field, groups):
    global row_ptr
    a.cell(row_ptr, 1, title).font = BB
    hdr = row_ptr + 1
    for i, h in enumerate([field.replace("_", " ").title(), "Employees", "Left", "Attrition rate"], 1):
        c = a.cell(hdr, i, h); c.font, c.fill, c.border = H, HF, BOX
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    r = hdr + 1
    first = r
    for g in groups:
        a.cell(r, 1, g)
        a.cell(r, 2, f'=COUNTIFS({rng(field)},$A{r})')
        a.cell(r, 3, f'=COUNTIFS({rng(field)},$A{r},{ATTR},"Yes")')
        a.cell(r, 4, f"=IF(B{r}=0,0,C{r}/B{r})").number_format = "0.0%"
        for i in range(1, 5):
            a.cell(r, i).font, a.cell(r, i).border = B, BOX
        r += 1
    blocks[title] = (hdr, first, r - 1)
    row_ptr = r + 2

block("Department", "department", ["Sales", "HR", "R&D"])
block("Age band", "age_band", ["Under 25", "25-34", "35-44", "45-54", "55+"])
block("Over time", "over_time", ["Yes", "No"])
block("Job role", "job_role", ["Sales Representative", "Laboratory Technician", "Human Resources", "Sales Executive",
                               "Research Scientist", "Manufacturing Director", "Healthcare Representative",
                               "Manager", "Research Director"])
block("Tenure (years at company)", "tenure_band", ["0-1 yr", "2-3 yrs", "4-5 yrs", "6-10 yrs", "10+ yrs"])

# overtime x marital status
a.cell(row_ptr, 1, "Overtime x marital status (attrition rate)").font = BB
hdr = row_ptr + 1
for i, h in enumerate(["Marital status", "No overtime", "Overtime", "Employees (overtime)"], 1):
    c = a.cell(hdr, i, h); c.font, c.fill, c.border = H, HF, BOX
    c.alignment = Alignment(horizontal="center", wrap_text=True)
for j, m in enumerate(["Single", "Married", "Divorced"], hdr + 1):
    a.cell(j, 1, m)
    for ccol, ot in ((2, "No"), (3, "Yes")):
        a.cell(j, ccol, f'=COUNTIFS({rng("marital_status")},$A{j},{rng("over_time")},"{ot}",{ATTR},"Yes")/'
                         f'COUNTIFS({rng("marital_status")},$A{j},{rng("over_time")},"{ot}")').number_format = "0.0%"
    a.cell(j, 4, f'=COUNTIFS({rng("marital_status")},$A{j},{rng("over_time")},"Yes")')
    for i in range(1, 5):
        a.cell(j, i).font, a.cell(j, i).border = B, BOX
marital_hdr = hdr
row_ptr = hdr + 6

def chart(title, key, anchor):
    hdr, first, lastr = blocks[key]
    ch = BarChart(); ch.type = "col"; ch.title = title
    ch.add_data(Reference(a, min_col=4, min_row=hdr, max_row=lastr), titles_from_data=True)
    ch.set_categories(Reference(a, min_col=1, min_row=first, max_row=lastr))
    ch.legend = None; ch.height, ch.width = 7, 13
    ch.y_axis.numFmt = "0%"; ch.y_axis.majorGridlines = None
    ch.y_axis.delete = False; ch.x_axis.delete = False
    ch.dataLabels = DataLabelList(); ch.dataLabels.showVal = True
    ch.dataLabels.showSerName = ch.dataLabels.showCatName = ch.dataLabels.showLegendKey = False
    ch.series[0].graphicalProperties.solidFill = "2B6CB0"
    a.add_chart(ch, anchor)

chart("Attrition rate by department", "Department", "F4")
chart("Attrition rate by age band", "Age band", "F19")
chart("Attrition rate by overtime", "Over time", "F34")
chart("Attrition rate by job role", "Job role", "F49")
chart("Attrition rate by tenure", "Tenure (years at company)", "F64")

wb.move_sheet("KPIs", offset=-1); wb.move_sheet("Analysis", offset=-1)   # KPIs, Analysis, Data
wb.active = 0
out = ROOT / "excel" / "HR_Attrition_Analysis.xlsx"
wb.save(out)
print("saved", out)
